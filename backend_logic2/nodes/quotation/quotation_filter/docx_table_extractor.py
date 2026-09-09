"""Word(.docx) 견적서에서 표를 직접 읽는 추출기.

이미지/PDF용 추출기(table_structure_extractor.py 등)와 근본적으로 다르다.
docx는 표 데이터가 이미지가 아니라 문서 안에 진짜 텍스트로 들어있어서 OCR이
전혀 필요 없다. python-docx로 표 구조(document.tables)를 그대로 읽기만 하면
되므로, 지금까지 봤던 인식 오류(오탈자, 행 뭉개짐, 손글씨 재확인 등)가
원천적으로 발생하지 않는다.

지원 형식: .docx만 지원한다. 구형 이진 형식 .doc는 python-docx가 읽지 못한다.
.doc 파일은 Word/한글/LibreOffice에서 열어 "다른 이름으로 저장 > Word 문서
(.docx)"로 변환한 뒤 사용하면 된다.

병합 셀 처리: python-docx는 병합된 셀을 rowspan/colspan으로 알려주지 않고,
"같은 셀 텍스트가 이웃 칸에도 그대로 반복"되는 형태로 보여준다. 이 스크립트는
그 반복을 그대로 둔다 — 값이 사라지거나 엉뚱한 칸으로 새는 이미지 추출기의
문제와는 다른 종류의(더 안전한) 결과다.

단독 실행 예:
    python -m backend_logic2.nodes.quotation.quotation_filter.docx_table_extractor `
      --docx "견적서.docx" --output-html docx_report.html
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path


def _cell_text(cell) -> str:
    """셀 자체 텍스트에 더해, 셀 안에 중첩된 표(표 안의 표)가 있으면 같이 풀어서 붙인다.

    docx는 이미지와 달리 "표 안의 표"가 문서 모델에 그대로 들어있어서 OCR처럼
    추측할 필요가 없다. cell.text만 읽으면 중첩 표 내용이 통째로 빠지므로
    cell.tables를 따로 순회해서 합친다. 2칸짜리 행(키: 값 형태)은 "키: 값"으로,
    그 외에는 칸을 띄어서 이어 붙인다.
    """
    parts = [cell.text.strip()] if cell.text.strip() else []
    for nested_table in cell.tables:
        for row in nested_table.rows:
            values = [c.text.strip() for c in row.cells]
            if len(values) == 2 and values[0] and values[1]:
                parts.append(f"{values[0]}: {values[1]}")
            else:
                parts.append(" ".join(v for v in values if v))
    return "; ".join(p for p in parts if p)


def extract_tables_from_docx(path: str | Path) -> list[list[list[str]]]:
    """docx 안의 모든 표를 행 x 열 텍스트 그리드로 읽는다."""
    from docx import Document

    suffix = Path(path).suffix.lower()
    if suffix != ".docx":
        raise ValueError(
            f"지원하지 않는 확장자입니다: {suffix or '(없음)'}. "
            ".doc(구형 이진 형식)는 Word/한글/LibreOffice에서 .docx로 변환한 뒤 사용하세요."
        )

    document = Document(str(path))
    tables: list[list[list[str]]] = []
    for table in document.tables:
        grid = [[_cell_text(cell) for cell in row.cells] for row in table.rows]
        tables.append(grid)
    return tables


def extract_paragraphs(path: str | Path) -> list[str]:
    """표 밖의 본문 문단(회사명, 비고 등)을 읽는다."""
    from docx import Document

    document = Document(str(path))
    return [p.text.strip() for p in document.paragraphs if p.text.strip()]


def build_html_report(tables: list[list[list[str]]], paragraphs: list[str], source_label: str) -> str:
    paragraphs_html = "".join(f"<p>{p}</p>" for p in paragraphs)

    tables_html = []
    for index, grid in enumerate(tables, start=1):
        rows_html = "".join(
            "<tr>" + "".join(f"<td>{cell}</td>" for cell in row) + "</tr>"
            for row in grid
        )
        tables_html.append(f"""
        <section class="table-block">
          <h3>표 {index} <span class="badge">{len(grid)}행 x {len(grid[0]) if grid else 0}열</span></h3>
          <div class="table-scroll"><table>{rows_html}</table></div>
        </section>
        """)

    return f"""<!doctype html>
<html lang="ko">
<head>
<meta charset="utf-8">
<title>DOCX 표 추출 결과</title>
<style>
  body {{ font-family: "Malgun Gothic", sans-serif; margin: 24px; background: #fafafa; color: #222; }}
  h1 {{ font-size: 20px; }}
  .notice {{ background: #e6f4ea; border: 1px solid #b7ddc3; padding: 10px 14px; border-radius: 6px; margin-bottom: 18px; }}
  .paragraphs {{ color: #555; font-size: 13px; margin-bottom: 20px; }}
  .table-block {{ margin-top: 24px; }}
  .badge {{ display: inline-block; font-size: 12px; background: #e6f4ea; color: #1e7d32; border-radius: 10px; padding: 2px 10px; margin-left: 8px; }}
  .table-scroll {{ overflow-x: auto; }}
  table {{ border-collapse: collapse; margin-top: 8px; }}
  td {{ border: 1px solid #999; padding: 6px 10px; font-size: 13px; white-space: nowrap; }}
</style>
</head>
<body>
  <h1>DOCX 표 추출 결과</h1>
  <div class="notice">
    이미지가 아니라 문서 안 텍스트를 직접 읽은 결과라 OCR 인식 오류가 없습니다.
    로컬 결과이며 ERPNext 등 외부 시스템에 저장되지 않았습니다.
  </div>
  <div class="paragraphs">원본: {source_label}<br>{paragraphs_html}</div>
  {"".join(tables_html)}
</body>
</html>
"""


def main() -> None:
    parser = argparse.ArgumentParser(description="Word(.docx) 견적서 표 추출")
    parser.add_argument("--docx", required=True, help=".docx 파일 경로")
    parser.add_argument("--output-json", help="표 그리드 JSON 저장 경로")
    parser.add_argument("--output-html", default="docx_report.html", help="HTML 리포트 저장 경로")
    args = parser.parse_args()

    tables = extract_tables_from_docx(args.docx)
    paragraphs = extract_paragraphs(args.docx)

    if args.output_json:
        Path(args.output_json).write_text(json.dumps(tables, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"표 그리드 JSON 저장: {args.output_json}")

    html = build_html_report(tables, paragraphs, source_label=args.docx)
    Path(args.output_html).write_text(html, encoding="utf-8")
    print(f"HTML 리포트 저장: {args.output_html}")
    print(f"찾은 표 개수: {len(tables)}")
    for index, grid in enumerate(tables, start=1):
        print(f"  표 {index}: {len(grid)}행 x {len(grid[0]) if grid else 0}열")


if __name__ == "__main__":
    main()
