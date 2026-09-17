from __future__ import annotations

import io
from decimal import Decimal
from email.message import EmailMessage
from unittest.mock import Mock

import pytest

from backend_logic2.nodes.quotation.quotation_filter.quotation_extractor import (
    DEFAULT_VISION_ADAPTER,
    DEFAULT_VISION_MODEL,
    FINETUNED_USER_PROMPT,
    LocalHuggingFaceQuotationParser,
    PreparedSource,
    SourceKind,
    VisionInput,
    _ParsedQuotation,
    _normalize_finetuned_quotation,
    _normalize_generated_quotation,
    _normalize_currency,
    apply_document_fallbacks,
    classify_source,
    extract_document_fallbacks,
    extract_quotation,
    extract_quotation_bytes,
    prepare_source,
)


def _parsed_quotation() -> _ParsedQuotation:
    return _ParsedQuotation.model_validate(
        {
            "quotation_id": "EST-001",
            "currency": "KRW",
            "subtotal": 1000,
            "tax_amount": 100,
            "total_amount": 1100,
            "items": [
                {
                    "item_name": "테스트 품목",
                    "quantity": 1,
                    "unit_price": 1000,
                    "amount": 1000,
                }
            ],
        }
    )


def test_finetuned_defaults_match_uploaded_adapter(monkeypatch) -> None:
    monkeypatch.setattr(
        "backend_logic2.nodes.quotation.quotation_filter.quotation_extractor._project_model_setting",
        lambda _name: None,
    )
    for name in (
        "HF_QUOTATION_VISION_MODEL",
        "HF_QUOTATION_VISION_ADAPTER",
        "HF_QUOTATION_VISION_MAX_PIXELS",
        "HF_QUOTATION_VISION_MAX_NEW_TOKENS",
    ):
        monkeypatch.delenv(name, raising=False)

    parser = LocalHuggingFaceQuotationParser()

    assert parser.vision_model_name == DEFAULT_VISION_MODEL == "Qwen/Qwen3.5-9B"
    assert parser.vision_adapter_name == DEFAULT_VISION_ADAPTER
    assert parser.vision_max_pixels == 802816
    assert parser.vision_max_new_tokens == 512
    assert "expected_delivery_date" in FINETUNED_USER_PROMPT
    assert "문서 하단" in FINETUNED_USER_PROMPT
    assert "유효기간" in FINETUNED_USER_PROMPT
    assert "확대 이미지" in FINETUNED_USER_PROMPT


def test_image_source_uses_direct_finetuned_extraction() -> None:
    parser = LocalHuggingFaceQuotationParser()
    expected = _parsed_quotation()
    parser._extract_images = Mock(return_value=expected)
    parser._structure_text = Mock(side_effect=AssertionError("text path must not run"))
    prepared = PreparedSource(
        kind=SourceKind.IMAGE,
        text="[견적서 이미지]",
        vision_inputs=[VisionInput(data=b"image", filename="quote.png")],
    )

    actual = parser(prepared, "RFQ-1", "공급사", [])

    assert actual is expected
    parser._extract_images.assert_called_once_with(prepared, [])


def test_text_source_keeps_existing_text_model_path() -> None:
    parser = LocalHuggingFaceQuotationParser()
    expected = _parsed_quotation()
    parser._structure_text = Mock(return_value=expected)
    parser._extract_images = Mock(side_effect=AssertionError("vision path must not run"))
    prepared = PreparedSource(kind=SourceKind.TEXT, text="견적 원문")

    actual = parser(prepared, "RFQ-1", "공급사", [])

    assert actual is expected
    parser._structure_text.assert_called_once()


def test_docx_source_preserves_paragraphs_and_table_structure(tmp_path) -> None:
    from docx import Document

    docx_path = tmp_path / "quote.docx"
    document = Document()
    document.add_paragraph("공급자: 테스트상사")
    table = document.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "품목명"
    table.rows[0].cells[1].text = "수량"
    table.rows[0].cells[2].text = "단가"
    table.rows[1].cells[0].text = "시약 A"
    table.rows[1].cells[1].text = "2"
    table.rows[1].cells[2].text = "1,000"
    document.save(docx_path)

    prepared = prepare_source(docx_path)

    assert classify_source(docx_path) == SourceKind.DOCX
    assert prepared.kind == SourceKind.DOCX
    assert prepared.vision_inputs == []
    assert "[본문]\n공급자: 테스트상사" in prepared.text
    assert "품목명 | 수량 | 단가" in prepared.text
    assert "시약 A | 2 | 1,000" in prepared.text
    assert any("OCR 미사용" in item for item in prepared.evidence)


def test_email_source_extracts_attached_docx_without_ocr(tmp_path) -> None:
    from docx import Document

    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("공급자: 첨부상사")
    document.save(buffer)

    message = EmailMessage()
    message["From"] = "vendor@example.com"
    message["To"] = "buyer@example.com"
    message["Subject"] = "견적서 송부"
    message.set_content("첨부 견적서를 확인해 주세요.")
    message.add_attachment(
        buffer.getvalue(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename="quote.docx",
    )
    email_path = tmp_path / "quote.eml"
    email_path.write_bytes(message.as_bytes())

    prepared = prepare_source(email_path)

    assert prepared.kind == SourceKind.EMAIL
    assert "[attachment: quote.docx]" in prepared.text
    assert "공급자: 첨부상사" in prepared.text
    assert any("OCR 미사용" in item for item in prepared.evidence)


def test_finetuned_output_normalizes_adapter_schema_without_recalculation() -> None:
    normalized = _normalize_finetuned_quotation(
        {
            "quotation_id": "EST-001",
            "supplier_name": "모델이 읽은 공급사",
            "currency": None,
            "subtotal": "1,000",
            "tax_amount": "100",
            # Deliberately inconsistent: image values must be preserved.
            "total_amount": "1,500",
            "items": [
                {
                    "item_code": "ITEM-1",
                    "item_name": "테스트 품목",
                    "quantity": "1",
                    "unit_price": "1,000",
                    "amount": "1,000",
                    "expected_delivery_date": "2026.09.20",
                }
            ],
        }
    )
    parsed = _ParsedQuotation.model_validate(normalized)

    assert parsed.supplier_name is None
    assert parsed.currency is None
    assert parsed.total_amount == Decimal("1500")
    assert parsed.items[0].expected_delivery_date.isoformat() == "2026-09-20"


def test_finetuned_output_infers_zero_tax_when_total_equals_subtotal() -> None:
    normalized = _normalize_finetuned_quotation(
        {
            "quotation_id": None,
            "currency": None,
            "subtotal": 3_900_000,
            "tax_amount": None,
            "total_amount": 3_900_000,
            "items": [
                {
                    "item_name": "ITEM-SUB-116",
                    "quantity": 1000,
                    "unit_price": 3900,
                    "amount": 3_900_000,
                }
            ],
        }
    )

    parsed = _ParsedQuotation.model_validate(normalized)

    assert parsed.tax_amount == Decimal("0")
    assert parsed.subtotal == parsed.total_amount == Decimal("3900000")


@pytest.mark.parametrize(("raw", "expected"), [
    ("원", "KRW"),
    ("$", "USD"),
    ("엔", "JPY"),
    ("€", "EUR"),
    ("RMB", "CNY"),
    (None, None),
    ("", None),
])
def test_currency_is_normalized_only_when_explicit(raw, expected) -> None:
    assert _normalize_currency(raw) == expected


def test_missing_currency_blocks_automatic_registration(tmp_path) -> None:
    image_path = tmp_path / "quote.png"
    image_path.write_bytes(b"test image bytes")

    def injected_parser(*_args):
        payload = _generated_payload_with_delivery(None)
        payload["currency"] = None
        return payload

    with pytest.raises(ValueError, match="결제 통화"):
        extract_quotation(
            image_path,
            "RFQ-001",
            supplier_name="공급사",
            model_parser=injected_parser,
        )


def test_expected_delivery_date_reaches_public_quotation_model(tmp_path) -> None:
    image_path = tmp_path / "quote.png"
    image_path.write_bytes(b"test image bytes")

    def injected_parser(*_args):
        return {
            "quotation_id": "EST-001",
            "currency": "KRW",
            "subtotal": 1000,
            "tax_amount": 100,
            "total_amount": 1100,
            "items": [
                {
                    "item_name": "테스트 품목",
                    "quantity": 1,
                    "unit_price": 1000,
                    "amount": 1000,
                    "expected_delivery_date": "2026-09-20",
                }
            ],
        }

    quotation = extract_quotation(
        image_path,
        "RFQ-001",
        supplier_name="공급사",
        model_parser=injected_parser,
    )

    assert quotation.items[0].expected_delivery_date.isoformat() == "2026-09-20"


def test_document_text_fills_omitted_terms_validity_and_delivery_date() -> None:
    text = """
    유 효 기 간 : 2026-09-30
    납 기 일 2026-09-
    30
    예 상 납 품 일 : 2026-09-30 ( 요 청 납 기 일 엄 수 )
    납 품 장 소 : 지 정 장 소 도 착 도
    사 양 특 기 : 전 용 하 드 케 이 스 포 함
    발 주 시 상 기 단 가 및 납 품 일 정 조 건 이 유 효 합 니 다.
    """
    fallbacks = extract_document_fallbacks(text)
    payload = {
        "valid_until": None,
        "notes": None,
        "items": [{"expected_delivery_date": None}],
    }

    apply_document_fallbacks(payload, fallbacks)

    assert payload["valid_until"] == "2026-09-30"
    assert payload["items"][0]["expected_delivery_date"] == "2026-09-30"
    assert "예상 납품일: 2026-09-30" in payload["notes"]
    assert "납품 장소:" in payload["notes"]
    assert "사양 특기:" in payload["notes"]


def test_document_fallbacks_do_not_replace_model_values() -> None:
    payload = {
        "valid_until": "2026-10-01",
        "notes": "모델이 추출한 특이사항",
        "items": [{"expected_delivery_date": "2026-10-02"}],
    }

    apply_document_fallbacks(payload, {
        "valid_until": "2026-09-30",
        "expected_delivery_date": "2026-09-30",
        "notes": "문서 보완값",
    })

    assert payload["valid_until"] == "2026-10-01"
    assert payload["notes"] == "모델이 추출한 특이사항"
    assert payload["items"][0]["expected_delivery_date"] == "2026-10-02"


def test_conflicting_validity_dates_are_not_silently_merged() -> None:
    fallbacks = extract_document_fallbacks(
        "유효기간: 2026-09-30\n유효기간: 2026-09-20"
    )
    payload = {"valid_until": None, "notes": None, "items": []}

    apply_document_fallbacks(payload, fallbacks)

    assert payload["valid_until"] is None
    assert fallbacks["conflicts"]["valid_until"] == [
        "2026-09-30",
        "2026-09-20",
    ]


def _generated_payload_with_delivery(
    delivery_date: str | None,
    lead_time_days: int | None = None,
) -> dict:
    return {
        "quotation_id": None,
        "currency": "KRW",
        "subtotal": 1000,
        "tax_amount": 100,
        "total_amount": 1100,
        "items": [{
            "item_code": "ITEM-1",
            "item_name": "테스트 품목",
            "quantity": 1,
            "unit_price": 1000,
            "amount": 1000,
            "expected_delivery_date": delivery_date,
            "lead_time_days": lead_time_days,
        }],
    }


def test_delivery_word_alone_does_not_authorize_model_date() -> None:
    normalized = _normalize_generated_quotation(
        _generated_payload_with_delivery("2026-09-30"),
        "배송 조건은 추후 협의합니다.",
        None,
    )

    assert normalized["items"][0]["expected_delivery_date"] is None


def test_model_delivery_value_must_match_explicit_document_evidence() -> None:
    matched = _normalize_generated_quotation(
        _generated_payload_with_delivery("2026-09-30", 14),
        "납품 예정일: 2026-09-30\n리드 타임: 2주",
        None,
    )
    mismatched = _normalize_generated_quotation(
        _generated_payload_with_delivery("2026-10-01", 10),
        "납품 예정일: 2026-09-30\n리드 타임: 2주",
        None,
    )

    assert matched["items"][0]["expected_delivery_date"] == "2026-09-30"
    assert matched["items"][0]["lead_time_days"] == 14
    assert mismatched["items"][0]["expected_delivery_date"] is None
    assert mismatched["items"][0]["lead_time_days"] is None


def test_conflicting_delivery_dates_are_not_silently_merged() -> None:
    fallbacks = extract_document_fallbacks(
        "납품 예정일: 2026-09-30\n출고 예정일: 2026-09-20"
    )
    payload = {
        "valid_until": None,
        "notes": None,
        "items": [{"expected_delivery_date": None}],
    }

    apply_document_fallbacks(payload, fallbacks)

    assert payload["items"][0]["expected_delivery_date"] is None
    assert fallbacks["conflicts"]["expected_delivery_date"] == [
        "2026-09-30",
        "2026-09-20",
    ]


def test_erp_attachment_bytes_are_extracted_without_local_path() -> None:
    def injected_parser(*_args):
        return {
            "quotation_id": None,
            "currency": "KRW",
            "subtotal": 1000,
            "tax_amount": 100,
            "total_amount": 1100,
            "items": [{
                "item_name": "테스트 품목",
                "quantity": 1,
                "unit_price": 1000,
                "amount": 1000,
            }],
        }

    quotation = extract_quotation_bytes(
        b"not-decoded-by-injected-parser",
        "quotation.jpg",
        "RFQ-001",
        supplier_name="공급사",
        supplier_id="SUP-001",
        fallback_quotation_id="EMAIL-COMM-1-FILE-1",
        message_id="COMM-1",
        content_type="image/jpeg",
        model_parser=injected_parser,
    )

    assert quotation.quotation_id == "EMAIL-COMM-1-FILE-1"
    assert quotation.source.path is None
    assert quotation.source.message_id == "COMM-1"
    assert quotation.source.content_type == "image/jpeg"
